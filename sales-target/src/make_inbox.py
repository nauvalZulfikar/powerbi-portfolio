"""Turn the clean AdventureWorks reseller extract into the messy Excel files a client would actually send.

    uv run python src/make_inbox.py      -> data/inbox/ (deterministic: same files every run)

Three branches, three formats:
- North America: a workbook per year, a sheet per month, title block, Month total row, the order
  date typed only on the first line of each order, a pasted
  duplicate sheet, a Notes sheet, rep names with stray case and spaces, the odd quantity as text.
- Europe: a workbook per quarter, German number and date text, "LASTNAME, First" reps, blank
  spacer rows, a Total row, an extra Comment column in one quarter.
- Pacific: a system export with "order/line" in one cell, Excel serial dates, Windows logins,
  and a resend that overlaps the first export by two months.
Plus the quarterly quota sheet ('000, wide, some numbers typed as text) and a reference workbook.
"""
import random
import shutil
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "source"
OUT = ROOT / "data" / "inbox"
RNG = random.Random(20260)

NICKNAMES = {281: "Mike Blythe", 286: "Pam Ansman-Wolfe", 288: "Jose Saraiva", 292: "Ranjit Varkey",
             272: "Steve Jiang", 289: "Dave Campbell"}
MANAGER_BRANCH = {272: "NA", 290: "EU", 294: "PAC"}


def read(name):
    with open(SOURCE / f"{name}.csv", encoding="utf-8") as f:
        return [line.rstrip("\n").split("|") for line in f if line.strip()]


def branch(territory_key):
    k = int(territory_key)
    return "NA" if 1 <= k <= 6 else "PAC" if k == 9 else "EU" if k in (7, 8, 10) else None


def german(x):
    s = f"{round(x, 2):,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def save(wb, path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"  {path.relative_to(OUT)}: {rows:,} rows")


def workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


# ------------------------------------------------------------------ source
def load():
    terr = {int(r[0]): {"region": r[2], "country": r[3], "group": r[4]} for r in read("DimSalesTerritory")}
    quota_reps = {int(r[1]) for r in read("FactSalesQuota")}
    emp = {int(r[0]): {"key": int(r[0]), "manager": r[1], "territory": r[4], "name": f"{r[5]} {r[6]}",
                       "first": r[5], "last": r[6], "title": r[9], "login": r[12]}
           for r in read("DimEmployee") if int(r[0]) in quota_reps}
    for e in emp.values():
        e["branch"] = MANAGER_BRANCH.get(e["key"]) or branch(e["territory"])
    geo = {r[0]: (r[1], r[5]) for r in read("DimGeography")}  # GeographyKey -> (City, Country)
    resellers = {int(r[0]): {"code": r[2], "type": r[4], "name": r[5], "city": geo[r[1]][0], "country": geo[r[1]][1]}
                 for r in read("DimReseller")}
    cats = {r[0]: r[2] for r in read("DimProductCategory")}
    subs = {r[0]: (r[2], cats[r[5]]) for r in read("DimProductSubcategory")}
    prod_rows = read("DimProduct")
    by_key = {int(r[0]): {"code": r[1], "name": r[5], "sub": r[2]} for r in prod_rows}
    latest = {}
    for r in sorted(prod_rows, key=lambda r: int(r[0])):
        if r[2] in subs:
            cost = float(r[8]) if r[8].strip() else (latest.get(r[1]) or {}).get("cost")  # StandardCost, latest version
            latest[r[1]] = {"code": r[1], "name": r[5], "sub": subs[r[2]][0], "cat": subs[r[2]][1], "cost": cost}
    lines = []
    for r in read("FactResellerSales"):
        p, s, e = by_key[int(r[0])], resellers[int(r[4])], emp[int(r[5])]
        lines.append({"so": r[9], "ln": int(r[10]), "qty": int(r[12]), "price": float(r[13]), "disc": float(r[15]),
                      "amount": float(r[19]), "date": datetime.strptime(r[24][:10], "%Y-%m-%d"),
                      "ship": datetime.strptime(r[26][:10], "%Y-%m-%d"),
                      "branch": branch(r[8]), "region": terr[int(r[8])]["region"], "country": terr[int(r[8])]["country"],
                      "product_code": p["code"], "product": p["name"], "reseller": s["name"], "reseller_code": s["code"],
                      "rep": e})
    lines.sort(key=lambda x: (x["date"], x["so"], x["ln"]))
    return terr, emp, resellers, latest, lines


# ------------------------------------------------------------------ branches
NA_HEADER = ["Order #", "Line", "Order Date", "Ship Date", "Reseller", "Sales Rep", "Territory", "Product Code", "Product", "Qty",
             "Unit Price", "Discount", "Customer No", "Line Total"]


def north_america(lines):
    years = defaultdict(lambda: defaultdict(list))
    for x in (x for x in lines if x["branch"] == "NA"):
        d = x["date"]
        years[max(d.year, 2011)][date(d.year, d.month, 1)].append(x)
    for year, months in sorted(years.items()):
        wb, n = workbook(), 0
        for month, xs in sorted(months.items()):
            name = month.strftime("%b %Y")
            body, prev = [], None
            for x in xs:
                rep = x["rep"]["name"]
                if RNG.random() < 0.05:
                    rep = rep.upper()
                elif RNG.random() < 0.05:
                    rep = rep.replace(" ", "  ", 1) + " "
                qty = str(x["qty"]) if RNG.random() < 0.03 else x["qty"]
                # The order date is typed once per order: later lines of the same order leave it blank.
                body.append([x["so"], x["ln"], x["date"] if x["so"] != prev else None, x["ship"], x["reseller"], rep, x["region"], x["product_code"], x["product"],
                             qty, x["price"], x["disc"], x["reseller_code"], x["amount"]])
                prev = x["so"]
            rows = [["Adventure Works - North America reseller sales"], [f"Month: {name}"], [], NA_HEADER, *body,
                    ["Month total"] + [None] * 12 + [sum(x["amount"] for x in xs)]]
            copies = [name, f"{name} (2)"] if name == "Mar 2012" else [name]
            for title in copies:
                ws = wb.create_sheet(title)
                for r in rows:
                    ws.append(r)
            n += len(body)
        wb.create_sheet("Notes")["A1"] = "Figures are preliminary until month close. Ask Stephen for corrections."
        save(wb, OUT / "sales" / "North America" / f"NA Reseller Sales {year}.xlsx", n)


EU_HEADER = ["Order", "Pos", "Date", "Customer", "Sales rep", "Country", "Item", "Qty", "Price", "Disc", "Net ", "Shipped"]


def europe(lines):
    quarters = defaultdict(list)
    for x in (x for x in lines if x["branch"] == "EU"):
        quarters[(x["date"].year, (x["date"].month - 1) // 3 + 1)].append(x)
    for (y, q), xs in sorted(quarters.items()):
        comment = (y, q) == (2012, 4)
        header = EU_HEADER + (["Comment"] if comment else [])
        wb = workbook()
        ws = wb.create_sheet("Sales")
        ws.append(header)
        for i, x in enumerate(xs, 1):
            e = x["rep"]
            row = [x["so"], x["ln"], x["date"].strftime("%d.%m.%Y"), x["reseller_code"], f"{e['last'].upper()}, {e['first']}",
                   x["country"], f"{x['product_code']} {x['product']}", x["qty"], german(x["price"]), x["disc"],
                   german(x["amount"]), x["ship"].strftime("%d.%m.%Y")]
            if comment:
                row.append("checked" if RNG.random() < 0.1 else None)
            ws.append(row)
            if i % 40 == 0:
                ws.append([""] * len(header))
        total = ["Total"] + [None] * (len(header) - 1)
        total[10] = german(sum(round(x["amount"], 2) for x in xs))
        ws.append(total)
        save(wb, OUT / "sales" / "Europe" / f"EU sales {y}-Q{q}.xlsx", len(xs))


def pacific(lines):
    xs = [x for x in lines if x["branch"] == "PAC"]
    for name, keep in [("Pacific sales export.xlsx", lambda d: d <= datetime(2013, 8, 31)),
                       ("Pacific sales export (1).xlsx", lambda d: d >= datetime(2013, 7, 1))]:
        part = [x for x in xs if keep(x["date"])]
        wb = workbook()
        ws = wb.create_sheet("Export")
        ws.append(["Order/Line", "Date", "Ship Date", "Customer", "Territory", "Rep", "SKU", "Product", "Qty", "Unit Price", "Discount",
                   "Amount"])
        for x in part:
            ws.append([f"{x['so']}/{x['ln']}", (x["date"].date() - date(1899, 12, 30)).days,
                       (x["ship"].date() - date(1899, 12, 30)).days, x["reseller_code"],
                       x["region"], x["rep"]["login"], x["product_code"], x["product"], x["qty"], x["price"], x["disc"],
                       x["amount"]])
        ws.append(["TOTAL"] + [None] * 10 + [sum(x["amount"] for x in part)])
        save(wb, OUT / "sales" / "Pacific" / name, len(part))


# ------------------------------------------------------------------ targets & reference
def targets(emp):
    rows = read("FactSalesQuota")
    keys = sorted({r[2] for r in rows})
    quarter = {k: f"{2011 + i // 4} Q{i % 4 + 1}" for i, k in enumerate(keys)}  # ordinal: source quarters are off
    quota = defaultdict(float)
    for r in rows:
        quota[(int(r[1]), quarter[r[2]])] += float(r[5])
    header, cols = ["Rep", "Region"], []
    for y in (2011, 2012, 2013):
        qs = [f"{y} Q{q}" for q in range(1, 5)]
        header += qs + [f"FY{y}"]
        cols.append(qs)
    wb = workbook()
    ws = wb.create_sheet("Quota")
    ws.append(["Sales quota by rep"])
    ws.append(["USD '000"])
    ws.append(header)
    totals, count = [0.0] * (len(header) - 2), 0
    for k in sorted(emp):
        row = [NICKNAMES.get(k, emp[k]["name"]), emp[k]["branch"]]
        for qs in cols:
            vals = [quota[(k, q)] / 1000 if (k, q) in quota else None for q in qs]
            for v in vals:
                count += 1
                row.append(f"{v:,.3f}" if v is not None and count % 7 == 0 else v)
            row.append(sum(v or 0 for v in vals))
        for i, v in enumerate(row[2:]):
            totals[i] += float(v.replace(",", "")) if isinstance(v, str) else v or 0
        ws.append(row)
    ws.append(["Total", None] + totals)
    save(wb, OUT / "targets" / "Sales quota 2011-2013.xlsx", len(emp))


def reference(terr, emp, resellers, products):
    wb = workbook()
    sheets = {
        "Reps": (["Rep Key", "Rep", "Login", "Title", "Territory Key", "Manager Key"],
                 [[k, e["name"], e["login"], e["title"], int(e["territory"]), int(e["manager"])] for k, e in sorted(emp.items())]),
        "Aliases": (["Alias", "Rep Key"], []),
        "Territories": (["Territory Key", "Region", "Country", "Group", "Branch"],
                        [[k, t["region"], t["country"], t["group"], branch(k)] for k, t in sorted(terr.items()) if k <= 10]),
        "Products": (["Product Code", "Product", "Subcategory", "Category", "Standard Cost"],
                     [[p["code"], p["name"], p["sub"], p["cat"], p["cost"]] for _, p in sorted(products.items())]),
        "Resellers": (["Reseller Key", "Reseller Code", "Reseller", "Business Type", "City", "Country"],
                      [[k, r["code"], r["name"], r["type"], r["city"], r["country"]] for k, r in sorted(resellers.items())]),
    }
    seen = set()
    for k, e in sorted(emp.items()):
        for alias in (e["name"], f"{e['last'].upper()}, {e['first']}", e["login"], NICKNAMES.get(k)):
            if alias and alias not in seen:
                seen.add(alias)
                sheets["Aliases"][1].append([alias, k])
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    save(wb, OUT / "reference" / "Reference.xlsx", sum(len(r) for _, r in sheets.values()))


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    terr, emp, resellers, products, lines = load()
    north_america(lines)
    europe(lines)
    pacific(lines)
    targets(emp)
    reference(terr, emp, resellers, products)


if __name__ == "__main__":
    main()
