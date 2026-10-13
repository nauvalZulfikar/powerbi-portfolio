"""Python mirror of the Power Query cleaning (powerquery/*.pq), for tests and the HTML twin.

    uv run python src/clean.py

Power BI runs the M; this script applies the same rules to the same inbox so the data tests can
prove the M logic reconciles to the source, and the preview has numbers to show. Every rule here
names the M query it mirrors. Writes data/marts/*.parquet.
"""
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import duckdb
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
INBOX = ROOT / "data" / "inbox"
MARTS = ROOT / "data" / "marts"
EXCEL_EPOCH = date(1899, 12, 30)


def key(v):
    """fnKey: trim, collapse repeated spaces, upper-case."""
    return None if v is None else " ".join(str(v).split()).upper()


def number(v):
    """Number.From: numbers pass, numeric text converts, anything else is null."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def integer(v):
    n = number(v)
    return int(n) if n is not None and n == int(n) else None


def german(v):
    """Number.FromText(_, "de-DE")."""
    if not isinstance(v, str) or not re.fullmatch(r"-?[\d.]*\d,?\d*", v.strip()):
        return None
    return float(v.strip().replace(".", "").replace(",", "."))


def as_date(v):
    """Date.From: datetimes and Excel serial numbers."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return EXCEL_EPOCH + timedelta(days=int(v))
    return None


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def promote(rows):
    header = [h if h is not None else f"Column{i + 1}" for i, h in enumerate(rows[0])]
    return [dict(zip(header, r + [None] * (len(header) - len(r)))) for r in rows[1:]]


def reference(sheet):
    """fnReference: one sheet of Reference.xlsx, headers promoted."""
    wb = openpyxl.load_workbook(INBOX / "reference" / "Reference.xlsx", read_only=True)
    return promote(sheet_rows(wb[sheet]))


# ------------------------------------------------------------- branch readers
def read_north_america(path):
    """fnReadNorthAmerica."""
    out = []
    for ws in openpyxl.load_workbook(path, read_only=True).worksheets:
        if ws.title == "Notes":
            continue
        last = None
        for i, r in enumerate(promote(sheet_rows(ws)[3:])):
            order = r["Order #"]
            last = r["Order Date"] if r["Order Date"] is not None else last  # Table.FillDown({"Order Date"})
            out.append({"sheet": ws.title, "row": 5 + i, "order_number": order, "order_line": r["Line"],
                        "order_date": as_date(last), "ship_date": as_date(r["Ship Date"]),
                        "reseller_code": r["Customer No"],
                        "rep_name": r["Sales Rep"], "territory": r["Territory"], "product_code": r["Product Code"],
                        "quantity": r["Qty"], "unit_price": r["Unit Price"], "discount": r["Discount"],
                        "amount": r["Line Total"],
                        "sheet_issue": "Duplicate sheet" if ws.title.endswith(")")
                        else "Subtotal row" if order == "Month total" else None})
    return out


def read_europe(path):
    """fnReadEurope."""
    out = []
    ws = openpyxl.load_workbook(path, read_only=True)["Sales"]
    def dmy(d):
        try:
            return datetime.strptime(d, "%d.%m.%Y").date() if isinstance(d, str) else None
        except ValueError:
            return None
    for i, r in enumerate(promote(sheet_rows(ws))):
        item = r["Item"]
        out.append({"sheet": "Sales", "row": 2 + i, "order_number": r["Order"], "order_line": r["Pos"],
                    "order_date": dmy(r["Date"]), "ship_date": dmy(r["Shipped"]), "reseller_code": r["Customer"], "rep_name": r["Sales rep"],
                    "territory": r["Country"], "product_code": None if item is None else str(item).split(" ")[0],
                    "quantity": r["Qty"], "unit_price": german(r["Price"]), "discount": r["Disc"],
                    "amount": german(r["Net "]),
                    "sheet_issue": "Subtotal row" if r["Order"] == "Total" else None})
    return out


def read_pacific(path):
    """fnReadPacific."""
    out = []
    ws = openpyxl.load_workbook(path, read_only=True)["Export"]
    for i, r in enumerate(promote(sheet_rows(ws))):
        parts = (str(r["Order/Line"]).split("/") if r["Order/Line"] is not None else [None]) + [None]
        out.append({"sheet": "Export", "row": 2 + i, "order_number": parts[0], "order_line": parts[1],
                    "order_date": as_date(r["Date"]), "ship_date": as_date(r["Ship Date"]),
                    "reseller_code": r["Customer"], "rep_name": r["Rep"],
                    "territory": r["Territory"], "product_code": r["SKU"], "quantity": r["Qty"],
                    "unit_price": r["Unit Price"], "discount": r["Discount"], "amount": r["Amount"],
                    "sheet_issue": "Subtotal row" if parts[0] == "TOTAL" else None})
    return out


READERS = {"North America": read_north_america, "Europe": read_europe, "Pacific": read_pacific}


# ------------------------------------------------------------------ queries
def staging():
    """Sales Staging: every row of every branch file, with the reason it can't be loaded."""
    aliases = {key(a["Alias"]): a["Rep Key"] for a in reference("Aliases")}
    products = {p["Product Code"] for p in reference("Products")}
    rows = []
    for branch, read in READERS.items():
        for path in sorted((INBOX / "sales" / branch).glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            for r in read(path):
                r.update(branch=branch, file=path.name)
                r["order_line"] = integer(r["order_line"])
                for k in ("quantity", "unit_price", "discount", "amount"):
                    r[k] = number(r[k])
                r["rep_key"] = aliases.get(key(r["rep_name"]))
                blank = r["order_number"] is None or str(r["order_number"]).strip() == ""
                r["issue"] = (r["sheet_issue"]
                              or ("Blank row" if blank else None)
                              or ("Bad value" if None in (r["order_date"], r["amount"], r["quantity"], r["order_line"]) else None)
                              or ("Unknown rep" if r["rep_key"] is None else None)
                              or ("Unknown product" if r["product_code"] not in products else None))
                rows.append(r)
    # A line sent twice: the first file by name wins.
    rows.sort(key=lambda r: (r["branch"], r["file"], r["sheet"], r["row"]))
    seen = set()
    for r in rows:
        if r["issue"] is None:
            k = (r["order_number"], r["order_line"])
            if k in seen:
                r["issue"] = "Duplicate line"
            seen.add(k)
    return rows


def targets():
    """Targets: quarterly quota per rep ('000, wide) -> one row per rep per day, in dollars."""
    ws = openpyxl.load_workbook(INBOX / "targets" / "Sales quota 2011-2013.xlsx", read_only=True)["Quota"]
    aliases = {key(a["Alias"]): a["Rep Key"] for a in reference("Aliases")}
    out = []
    for r in promote(sheet_rows(ws)[2:]):
        if r["Rep"] in (None, "Total"):
            continue
        for col, v in r.items():
            if not (isinstance(col, str) and len(col) == 7 and col[4:6] == " Q") or v is None:
                continue
            quota = 1000 * (float(v.replace(",", "")) if isinstance(v, str) else v)
            start = date(int(col[:4]), 3 * int(col[-1]) - 2, 1)
            end = date(start.year + (start.month == 10), (start.month + 2) % 12 + 1, 1)
            days = (end - start).days
            out += [{"rep_key": aliases.get(key(r["Rep"])), "date": start + timedelta(d), "quarter": col,
                     "daily_quota": quota / days} for d in range(days)]
    return out


def category_plan_sql():
    """Category Plan (SYNTHETIC): each fiscal quarter's quota up to the last sale, split by the prior fiscal
    year's category mix (the first fiscal year uses its own). Fiscal year starts in July."""
    fy = "(year({d}) + (month({d}) >= 7)::int)"
    fq = "'FY' || " + fy + " || ' Q' || ((((month({d}) + 5) % 12)) // 3 + 1)"
    return f"""create table category_plan as
        with quota as (select {fq.format(d='date')} fiscal_quarter, {fy.format(d='date')} fy, sum(daily_quota) quota
                       from targets where date <= (select max(order_date) from sales) group by all),
             by_fy as (select {fy.format(d='s.order_date')} fy, p.category, sum(s.amount) v
                       from sales s join products p using (product_code) group by all),
             mix as (select fy, category, v / sum(v) over (partition by fy) as share from by_fy),
             q as (select *, case when fy - 1 in (select fy from mix) then fy - 1 else (select min(fy) from mix) end basis
                   from quota)
        select q.fiscal_quarter, m.category, q.quota * m.share plan_amount,
               case when q.basis = q.fy - 1 then 'Prior fiscal year mix' else 'Same fiscal year mix (no history)' end mix_basis
        from q join mix m on m.fy = q.basis order by 1, 2"""


def main():
    MARTS.mkdir(parents=True, exist_ok=True)
    rows = staging()
    c = duckdb.connect()
    cols = ["branch", "file", "sheet", "row", "order_number", "order_line", "order_date", "ship_date", "reseller_code", "rep_name",
            "rep_key", "territory", "product_code", "quantity", "unit_price", "discount", "amount", "issue"]
    c.execute("""create table staging (branch varchar, file varchar, sheet varchar, row int, order_number varchar,
                 order_line int, order_date date, ship_date date, reseller_code varchar, rep_name varchar, rep_key int, territory varchar,
                 product_code varchar, quantity double, unit_price double, discount double, amount double, issue varchar)""")
    c.executemany(f"insert into staging values ({', '.join('?' * len(cols))})",
                  [[None if r[k] is None else str(r[k]) if k in ("order_number", "reseller_code", "rep_name") else r[k]
                    for k in cols] for r in rows])
    for sheet, name in [("Reps", "reps"), ("Territories", "territories"), ("Products", "products"), ("Resellers", "resellers")]:
        data = reference(sheet)
        names = [re.sub(r"\W+", "_", k.lower()) for k in data[0]]
        c.execute(f"create table {name} ({', '.join(f'"{n}" varchar' for n in names)})")
        c.executemany(f"insert into {name} values ({', '.join('?' * len(names))})",
                      [[None if v is None else str(v) for v in d.values()] for d in data])
    c.execute("""create table targets (rep_key int, date date, quarter varchar, daily_quota double)""")
    c.executemany("insert into targets values (?, ?, ?, ?)", [list(t.values()) for t in targets()])
    c.execute("""create table sales as
                 select s.order_number, s.order_line, s.order_date, s.ship_date, s.reseller_code, s.rep_key,
                        t.territory_key::int territory_key, s.product_code, s.quantity::int quantity,
                        s.unit_price, s.discount, s.amount, s.branch
                 from staging s left join territories t on t.region = s.territory
                 where s.issue is null""")
    c.execute("""create table cleaning_log as
                 select branch, file, coalesce(issue, 'Loaded') outcome, count(*) "rows", sum(amount) amount
                 from staging group by all order by all""")
    c.execute(category_plan_sql())
    for t in ("staging", "sales", "cleaning_log", "targets", "reps", "territories", "products", "resellers",
              "category_plan"):
        c.execute(f"copy {t} to '{MARTS / t}.parquet' (format parquet)")
        print(f"  {t:<13} {c.execute(f'select count(*) from {t}').fetchone()[0]:>8,} rows")


if __name__ == "__main__":
    main()
